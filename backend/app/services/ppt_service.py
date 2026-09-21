"""
PowerPoint Service

This service extracts structural and visual information
from PowerPoint presentations.

Responsibilities:
    - Read presentation metadata
    - Analyze slides
    - Extract text
    - Detect layouts
    - Extract fonts
    - Extract colors
    - Extract shape positions and dimensions
    - Detect images, tables, and charts
    - Detect slide background information

The extracted information is later used by the
PPT Analysis Agent to understand the uploaded template.

Important:
    python-pptx primarily supports .pptx files.
    Legacy .ppt files should be converted to .pptx
    before detailed analysis.
"""

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


# -------------------------------------------------------------------
# Helper Functions
# -------------------------------------------------------------------

def _emu_to_inches(value: int | None) -> float | None:
    """
    Convert PowerPoint EMU units into inches.

    PowerPoint internally stores positions and dimensions
    using English Metric Units (EMU).

    914400 EMU = 1 inch.
    """
    if value is None:
        return None

    return round(value / 914400, 2)


def _extract_color(color_format) -> dict[str, Any] | None:
    """
    Safely extract color information.

    PowerPoint colors can be:
        - RGB colors
        - Theme colors
        - Automatic colors

    Not every shape contains every color property, so
    this function intentionally handles missing values.
    """

    if color_format is None:
        return None

    color_info: dict[str, Any] = {}

    try:
        if color_format.type is not None:
            color_info["type"] = str(color_format.type)
    except Exception:
        pass

    # Try to extract RGB color.
    try:
        if color_format.rgb:
            color_info["rgb"] = str(color_format.rgb)
    except Exception:
        pass

    # Try to extract theme color.
    try:
        if color_format.theme_color:
            color_info["theme_color"] = str(color_format.theme_color)
    except Exception:
        pass

    # Try to extract brightness.
    try:
        if color_format.brightness is not None:
            color_info["brightness"] = color_format.brightness
    except Exception:
        pass

    return color_info if color_info else None


def _extract_font_information(shape) -> list[dict[str, Any]]:
    """
    Extract font information from text-containing shapes.
    """

    fonts = []

    if not hasattr(shape, "text_frame"):
        return fonts

    if not shape.has_text_frame:
        return fonts

    for paragraph in shape.text_frame.paragraphs:

        for run in paragraph.runs:

            font = run.font

            fonts.append(
                {
                    "name": font.name,
                    "size": font.size.pt if font.size else None,
                    "bold": font.bold,
                    "italic": font.italic,
                    "underline": font.underline,
                    "color": (
                        _extract_color(font.color)
                        if font.color
                        else None
                    ),
                }
            )

    return fonts


def _extract_text(shape) -> str:
    """
    Extract visible text from a shape.
    """

    if not hasattr(shape, "text"):
        return ""

    try:
        return shape.text.strip()
    except Exception:
        return ""


def _detect_shape_type(shape) -> str:
    """
    Convert python-pptx shape types into readable labels.
    """

    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "image"

    if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
        return "table"

    if shape.shape_type == MSO_SHAPE_TYPE.CHART:
        return "chart"

    if hasattr(shape, "text_frame") and shape.has_text_frame:
        return "text"

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"

    return "other"


def _extract_shape_position(shape) -> dict[str, float | None]:
    """
    Extract the position and size of a shape.

    These values help the generator understand
    where elements are placed inside the template.
    """

    return {
        "left": _emu_to_inches(shape.left),
        "top": _emu_to_inches(shape.top),
        "width": _emu_to_inches(shape.width),
        "height": _emu_to_inches(shape.height),
    }


def _extract_fill_information(shape) -> dict[str, Any] | None:
    """
    Extract basic fill/background information from a shape.
    """

    if not hasattr(shape, "fill"):
        return None

    try:
        fill = shape.fill

        fill_info: dict[str, Any] = {
            "type": str(fill.type) if fill.type else None
        }

        try:
            fill_info["color"] = _extract_color(fill.fore_color)
        except Exception:
            fill_info["color"] = None

        return fill_info

    except Exception:
        return None


def _extract_line_information(shape) -> dict[str, Any] | None:
    """
    Extract basic border/line styling.
    """

    if not hasattr(shape, "line"):
        return None

    try:
        line = shape.line

        return {
            "color": _extract_color(line.color),
            "width": (
                _emu_to_inches(line.width)
                if line.width is not None
                else None
            ),
        }

    except Exception:
        return None


def _extract_shape_information(shape) -> dict[str, Any]:
    """
    Extract all useful information from one PowerPoint shape.
    """

    shape_type = _detect_shape_type(shape)

    information = {
        "shape_type": shape_type,
        "name": getattr(shape, "name", None),
        "text": _extract_text(shape),
        "position": _extract_shape_position(shape),
        "fill": _extract_fill_information(shape),
        "line": _extract_line_information(shape),
        "fonts": _extract_font_information(shape),
    }

    return information


def _extract_slide_background(slide) -> dict[str, Any] | None:
    """
    Extract basic slide background information.

    Not every presentation uses a directly accessible
    solid background, so failures are handled safely.
    """

    try:
        background = slide.background
        fill = background.fill

        background_info = {
            "type": str(fill.type) if fill.type else None,
            "color": None,
        }

        try:
            background_info["color"] = _extract_color(
                fill.fore_color
            )
        except Exception:
            pass

        return background_info

    except Exception:
        return None


# -------------------------------------------------------------------
# Main PPT Extraction Function
# -------------------------------------------------------------------

def extract_ppt_structure(file_path: str) -> dict[str, Any]:
    """
    Extract structural and visual information from a PPTX file.

    Parameters:
        file_path:
            Path to the uploaded PPTX file.

    Returns:
        A structured dictionary containing:
            - presentation metadata
            - slide dimensions
            - slide layouts
            - slide text
            - fonts
            - colors
            - shape positions
            - images
            - tables
            - charts
            - background information
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"Presentation not found: {file_path}"
        )

    # python-pptx works with PPTX.
    # Legacy .ppt files require conversion first.
    if path.suffix.lower() == ".ppt":
        raise ValueError(
            "Legacy .ppt files are not directly supported. "
            "Please convert the file to .pptx before analysis."
        )

    presentation = Presentation(file_path)

    slides = []

    # ---------------------------------------------------------------
    # Process every slide
    # ---------------------------------------------------------------

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_text = []
        fonts = []
        shapes = []
        shape_types = []

        layout_name = slide.slide_layout.name

        # -----------------------------------------------------------
        # Process every shape on the slide
        # -----------------------------------------------------------

        for shape in slide.shapes:

            shape_type = _detect_shape_type(shape)

            shape_types.append(shape_type)

            text = _extract_text(shape)

            if text:
                slide_text.append(text)

            fonts.extend(
                _extract_font_information(shape)
            )

            shapes.append(
                _extract_shape_information(shape)
            )

        # -----------------------------------------------------------
        # Extract slide title
        # -----------------------------------------------------------

        title = ""

        try:
            if slide.shapes.title:
                title = slide.shapes.title.text.strip()
        except Exception:
            title = ""

        # -----------------------------------------------------------
        # Store complete slide information
        # -----------------------------------------------------------

        slides.append(
            {
                "slide_number": slide_number,
                "title": title,
                "layout": layout_name,
                "text": slide_text,
                "shape_types": shape_types,
                "shapes": shapes,
                "fonts": fonts,

                "has_image": "image" in shape_types,
                "has_table": "table" in shape_types,
                "has_chart": "chart" in shape_types,

                "background": _extract_slide_background(slide),
            }
        )

    # ---------------------------------------------------------------
    # Presentation-level information
    # ---------------------------------------------------------------

    return {
        "file_type": "pptx",
        "filename": path.name,

        "slide_count": len(slides),

        # Dimensions are kept in EMU for exact PowerPoint compatibility.
        "slide_width": presentation.slide_width,
        "slide_height": presentation.slide_height,

        # Human-readable dimensions.
        "slide_width_inches": _emu_to_inches(
            presentation.slide_width
        ),
        "slide_height_inches": _emu_to_inches(
            presentation.slide_height
        ),

        "slides": slides,
    }