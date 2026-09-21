"""
Document Service

This service extracts text and useful metadata from supported
document formats.

Supported formats:

    PDF
    DOCX
    PPTX
    PNG
    JPG
    JPEG
    WEBP

For scanned PDFs and images, OCR is used when normal text
extraction is not available.

The extracted text becomes the input for:

    Document Analysis
    RAG
    Embeddings
    Search
    AI Generation
"""

from pathlib import Path
from typing import Any

import pymupdf
from docx import Document
from pptx import Presentation

from app.services.ocr_service import (
    extract_text_from_image,
    extract_text_from_image_object,
)


# ===============================================================
# PDF EXTRACTION
# ===============================================================
def extract_text_from_pdf(
    file_path: str,
) -> str:
    """
    Extract text from a PDF.

    If a page contains no selectable text, the page is rendered
    as an image and OCR is applied.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"PDF file not found: {file_path}"
        )

    document = pymupdf.open(path)

    page_texts: list[str] = []

    try:
        for page_number, page in enumerate(document):

            text = page.get_text("text").strip()

            if text:
                page_texts.append(text)
                continue

            pixmap = page.get_pixmap(
                matrix=pymupdf.Matrix(2, 2),
                alpha=False,
            )

            image_bytes = pixmap.tobytes("png")

            import io
            from PIL import Image

            image = Image.open(
                io.BytesIO(image_bytes)
            )

            ocr_text = extract_text_from_image_object(
                image
            )

            if ocr_text:
                page_texts.append(
                    f"[Page {page_number + 1}]\n{ocr_text}"
                )

    finally:
        document.close()

    return "\n\n".join(page_texts).strip()
# ===============================================================
# DOCX EXTRACTION
# ===============================================================

def extract_text_from_docx(
    file_path: str,
) -> str:
    """
    Extract paragraphs and table content from a DOCX file.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"DOCX file not found: {file_path}"
        )

    document = Document(path)

    extracted_parts: list[str] = []

    # -----------------------------------------------------------
    # Paragraphs
    # -----------------------------------------------------------

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            extracted_parts.append(text)

    # -----------------------------------------------------------
    # Tables
    # -----------------------------------------------------------

    for table_index, table in enumerate(
        document.tables,
        start=1,
    ):

        extracted_parts.append(
            f"\n[TABLE {table_index}]"
        )

        for row in table.rows:

            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            extracted_parts.append(
                " | ".join(cells)
            )

    return "\n".join(extracted_parts).strip()


# ===============================================================
# PPTX EXTRACTION
# ===============================================================

def extract_text_from_pptx(
    file_path: str,
) -> str:
    """
    Extract text from PowerPoint slides.

    Text from every slide and text-containing shape is collected.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"PPTX file not found: {file_path}"
        )

    presentation = Presentation(path)

    slide_texts: list[str] = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):

        slide_parts: list[str] = [
            f"[SLIDE {slide_number}]"
        ]

        for shape in slide.shapes:

            # ---------------------------------------------------
            # Normal text shapes
            # ---------------------------------------------------

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    slide_parts.append(text)

            # ---------------------------------------------------
            # Tables
            # ---------------------------------------------------

            if getattr(shape, "has_table", False):

                for row in shape.table.rows:

                    cells = [
                        cell.text.strip()
                        for cell in row.cells
                    ]

                    slide_parts.append(
                        " | ".join(cells)
                    )

        if len(slide_parts) > 1:
            slide_texts.append(
                "\n".join(slide_parts)
            )

    return "\n\n".join(slide_texts).strip()


# ===============================================================
# IMAGE EXTRACTION
# ===============================================================

def extract_text_from_image_file(
    file_path: str,
) -> str:
    """
    Extract text from an image using OCR.
    """

    return extract_text_from_image(
        file_path
    )


# ===============================================================
# MAIN DOCUMENT DISPATCHER
# ===============================================================

def extract_document(
    file_path: str,
) -> dict[str, Any]:
    """
    Extract text from a supported document.

    This is the main function used by the API and RAG pipeline.

    Example:

        result = extract_document(
            "storage/uploads/document.pdf"
        )

        text = result["text"]
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"Document not found: {file_path}"
        )

    extension = path.suffix.lower()

    # -----------------------------------------------------------
    # PDF
    # -----------------------------------------------------------

    if extension == ".pdf":

        text = extract_text_from_pdf(
            str(path)
        )

    # -----------------------------------------------------------
    # DOCX
    # -----------------------------------------------------------

    elif extension == ".docx":

        text = extract_text_from_docx(
            str(path)
        )

    # -----------------------------------------------------------
    # PPTX
    # -----------------------------------------------------------

    elif extension == ".pptx":

        text = extract_text_from_pptx(
            str(path)
        )

    # -----------------------------------------------------------
    # Images
    # -----------------------------------------------------------

    elif extension in {
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
    }:

        text = extract_text_from_image_file(
            str(path)
        )

    # -----------------------------------------------------------
    # Legacy PPT
    # -----------------------------------------------------------

    elif extension == ".ppt":

        raise ValueError(
            "Legacy .ppt files are not directly supported by "
            "python-pptx. Please convert the file to .pptx."
        )

    # -----------------------------------------------------------
    # Unsupported format
    # -----------------------------------------------------------

    else:

        raise ValueError(
            f"Unsupported document type: {extension}"
        )

    return {
        "filename": path.name,
        "file_type": extension.replace(".", ""),
        "text": text,
        "character_count": len(text),
        "word_count": len(text.split()),
    }