"""
Document Generator Agent

Generates editable Microsoft Word documents.

Capabilities:
    - Creates DOCX documents
    - Generates title and heading sections
    - Supports paragraphs
    - Supports bullet points
    - Supports numbered lists
    - Supports tables
    - Supports an existing DOCX template
    - Preserves basic template styles
    - Saves generated files to storage/generated/
"""

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Pt

from app.config import settings


class DocumentGeneratorAgent:
    """
    Agent responsible for generating editable DOCX documents.
    """

    def __init__(self):

        self.output_directory = Path(
            settings.generated_directory
        )

        self.output_directory.mkdir(
            parents=True,
            exist_ok=True,
        )

    # ===========================================================
    # PUBLIC GENERATION METHOD
    # ===========================================================

    def generate(
        self,
        content: dict[str, Any],
        filename: str = "generated_document.docx",
        template_path: str | None = None,
    ) -> dict[str, Any]:
        """
        Generate an editable DOCX document.

        Parameters
        ----------
        content:
            Structured document content.

        filename:
            Output filename.

        template_path:
            Optional DOCX template.

        Returns
        -------
        dict
            Generation result.
        """

        if not isinstance(content, dict):

            raise ValueError(
                "Document content must be a dictionary."
            )

        # -------------------------------------------------------
        # Load template or create new document
        # -------------------------------------------------------

        template_used = False

        if template_path:

            template = Path(
                template_path
            )

            if not template.exists():

                raise FileNotFoundError(
                    f"Template not found: {template}"
                )

            if template.suffix.lower() != ".docx":

                raise ValueError(
                    "Only .docx templates are supported."
                )

            document = Document(
                str(template)
            )

            template_used = True

        else:

            document = Document()

        # -------------------------------------------------------
        # Filename
        # -------------------------------------------------------

        filename = Path(
            filename
        ).name

        if not filename.lower().endswith(
            ".docx"
        ):

            filename += ".docx"

        output_path = (
            self.output_directory
            / filename
        )

        # -------------------------------------------------------
        # Generate title
        # -------------------------------------------------------

        title = content.get(
            "title",
            "Generated Document",
        )

        self._add_title(
            document,
            str(title),
        )

        # -------------------------------------------------------
        # Subtitle
        # -------------------------------------------------------

        subtitle = content.get(
            "subtitle",
            "",
        )

        if subtitle:

            self._add_subtitle(
                document,
                str(subtitle),
            )

        # -------------------------------------------------------
        # Sections
        # -------------------------------------------------------

        sections = content.get(
            "sections",
            [],
        )

        for section in sections:

            if not isinstance(
                section,
                dict,
            ):
                continue

            self._add_section(
                document,
                section,
            )

        # -------------------------------------------------------
        # Simple paragraphs
        # -------------------------------------------------------

        paragraphs = content.get(
            "paragraphs",
            [],
        )

        for paragraph in paragraphs:

            if paragraph:

                document.add_paragraph(
                    str(paragraph)
                )

        # -------------------------------------------------------
        # Tables
        # -------------------------------------------------------

        tables = content.get(
            "tables",
            [],
        )

        for table_data in tables:

            if isinstance(
                table_data,
                dict,
            ):

                self._add_table(
                    document,
                    table_data,
                )

        # -------------------------------------------------------
        # Save
        # -------------------------------------------------------

        document.save(
            str(output_path)
        )

        # -------------------------------------------------------
        # Validate
        # -------------------------------------------------------

        if not output_path.exists():

            raise RuntimeError(
                "DOCX file was not created."
            )

        if output_path.stat().st_size == 0:

            raise RuntimeError(
                "Generated DOCX file is empty."
            )

        # -------------------------------------------------------
        # Return result
        # -------------------------------------------------------

        return {

            "success": True,

            "filename": output_path.name,

            "file_path": str(
                output_path
            ),

            "file_type": "docx",

            "editable": True,

            "template_used": template_used,

            "paragraph_count": len(
                document.paragraphs
            ),

            "table_count": len(
                document.tables
            ),
        }

    # ===========================================================
    # TITLE
    # ===========================================================

    def _add_title(
        self,
        document: Document,
        title: str,
    ) -> None:
        """
        Add document title.
        """

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            title
        )

        run.bold = True
        run.font.size = Pt(24)

    # ===========================================================
    # SUBTITLE
    # ===========================================================

    def _add_subtitle(
        self,
        document: Document,
        subtitle: str,
    ) -> None:
        """
        Add document subtitle.
        """

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            subtitle
        )

        run.italic = True
        run.font.size = Pt(14)

    # ===========================================================
    # SECTION
    # ===========================================================

    def _add_section(
        self,
        document: Document,
        section: dict[str, Any],
    ) -> None:
        """
        Add a structured document section.
        """

        heading = section.get(
            "heading",
            "",
        )

        if heading:

            document.add_heading(
                str(heading),
                level=1,
            )

        content = section.get(
            "content",
            "",
        )

        if content:

            document.add_paragraph(
                str(content)
            )

        bullets = section.get(
            "bullets",
            [],
        )

        for bullet in bullets:

            if bullet:

                paragraph = document.add_paragraph(
                    str(bullet),
                    style="List Bullet",
                )

        numbered = section.get(
            "numbered",
            [],
        )

        for item in numbered:

            if item:

                document.add_paragraph(
                    str(item),
                    style="List Number",
                )

    # ===========================================================
    # TABLE
    # ===========================================================

    def _add_table(
        self,
        document: Document,
        table_data: dict[str, Any],
    ) -> None:
        """
        Add a table to the document.

        Expected structure:

        {
            "headers": ["Technology", "Purpose"],
            "rows": [
                ["FastAPI", "Backend API"],
                ["ChromaDB", "Vector Database"]
            ]
        }
        """

        headers = table_data.get(
            "headers",
            [],
        )

        rows = table_data.get(
            "rows",
            [],
        )

        if not headers:

            return

        table = document.add_table(
            rows=1,
            cols=len(headers),
        )

        table.style = "Table Grid"

        # -------------------------------------------------------
        # Header
        # -------------------------------------------------------

        header_cells = table.rows[0].cells

        for index, header in enumerate(
            headers
        ):

            header_cells[index].text = str(
                header
            )

            for paragraph in (
                header_cells[index]
                .paragraphs
            ):

                for run in paragraph.runs:

                    run.bold = True

        # -------------------------------------------------------
        # Rows
        # -------------------------------------------------------

        for row in rows:

            cells = table.add_row().cells

            for index, value in enumerate(
                row
            ):

                if index < len(cells):

                    cells[index].text = str(
                        value
                    )

    # ===========================================================
    # ADD PAGE BREAK
    # ===========================================================

    def add_page_break(
        self,
        document: Document,
    ) -> None:
        """
        Add a page break.
        """

        document.add_page_break()