"""
Validation Agent

Responsible for validating generated DOCX and PPTX artifacts.

Validation includes:
    - File existence
    - File size
    - File extension
    - Basic file integrity
    - DOCX readability
    - PPTX readability
"""

from pathlib import Path
from typing import Any


class ValidationAgent:
    """
    Agent responsible for validating generated artifacts.
    """

    # ==========================================================
    # INITIALIZATION
    # ==========================================================

    def __init__(self):
        """
        Initialize the validation agent.
        """

        self.supported_extensions = {
            ".docx",
            ".pptx",
            ".pdf",
        }

    # ==========================================================
    # VALIDATE SINGLE FILE
    # ==========================================================

    def validate_file(
        self,
        file_path: str | Path,
    ) -> dict[str, Any]:
        """
        Validate a single generated file.

        Parameters
        ----------
        file_path:
            Path to the file that should be validated.

        Returns
        -------
        dict
            Validation result.
        """

        path = Path(file_path)

        result: dict[str, Any] = {
            "file_path": str(path),
            "filename": path.name,
            "exists": False,
            "size": 0,
            "extension": path.suffix.lower(),
            "supported": False,
            "readable": False,
            "valid": False,
            "errors": [],
        }

        # ------------------------------------------------------
        # Check existence
        # ------------------------------------------------------

        if not path.exists():

            result["errors"].append(
                "File does not exist."
            )

            return result

        result["exists"] = True

        # ------------------------------------------------------
        # Check file size
        # ------------------------------------------------------

        try:

            result["size"] = path.stat().st_size

        except OSError as exc:

            result["errors"].append(
                f"Unable to read file size: {exc}"
            )

            return result

        if result["size"] <= 0:

            result["errors"].append(
                "File is empty."
            )

            return result

        # ------------------------------------------------------
        # Check extension
        # ------------------------------------------------------

        if result["extension"] not in (
            self.supported_extensions
        ):

            result["errors"].append(
                f"Unsupported file type: "
                f"{result['extension']}"
            )

            return result

        result["supported"] = True

        # ------------------------------------------------------
        # Validate actual file structure
        # ------------------------------------------------------

        try:

            if result["extension"] == ".docx":

                self._validate_docx(path)

            elif result["extension"] == ".pptx":

                self._validate_pptx(path)

            elif result["extension"] == ".pdf":

                self._validate_pdf(path)

            result["readable"] = True

        except Exception as exc:

            result["errors"].append(
                f"File integrity validation failed: {exc}"
            )

            return result

        # ------------------------------------------------------
        # Final validation status
        # ------------------------------------------------------

        result["valid"] = (
            result["exists"]
            and result["size"] > 0
            and result["supported"]
            and result["readable"]
            and len(result["errors"]) == 0
        )

        return result

    # ==========================================================
    # DOCX VALIDATION
    # ==========================================================

    def _validate_docx(
        self,
        path: Path,
    ) -> None:
        """
        Validate that a DOCX file can be opened.
        """

        from docx import Document

        document = Document(
            str(path)
        )

        # Accessing these collections ensures
        # the document structure can be parsed.

        _ = document.paragraphs
        _ = document.tables

    # ==========================================================
    # PPTX VALIDATION
    # ==========================================================

    def _validate_pptx(
        self,
        path: Path,
    ) -> None:
        """
        Validate that a PPTX file can be opened.
        """

        from pptx import Presentation

        presentation = Presentation(
            str(path)
        )

        # Access slides to ensure the
        # presentation structure is readable.

        _ = presentation.slides

    # ==========================================================
    # PDF VALIDATION
    # ==========================================================

    def _validate_pdf(
        self,
        path: Path,
    ) -> None:
        """
        Validate that a PDF can be opened.
        """

        import pymupdf

        document = pymupdf.open(
            str(path)
        )

        # Access page count.

        _ = len(document)

        document.close()

    # ==========================================================
    # VALIDATE MULTIPLE ARTIFACTS
    # ==========================================================

    def validate_artifacts(
        self,
        artifacts: list[str | Path],
    ) -> dict[str, Any]:
        """
        Validate multiple generated artifacts.

        Parameters
        ----------
        artifacts:
            List of file paths.

        Returns
        -------
        dict
            Combined validation result.
        """

        results = []

        for artifact in artifacts:

            result = self.validate_file(
                artifact
            )

            results.append(result)

        all_valid = (
            len(results) > 0
            and all(
                item["valid"]
                for item in results
            )
        )

        return {
            "valid": all_valid,
            "artifact_count": len(results),
            "artifacts": results,
        }

    # ==========================================================
    # VALIDATE WORKFLOW OUTPUT
    # ==========================================================

    def validate_workflow_output(
        self,
        workflow_result: dict[str, Any],
    ) -> dict[str, Any]:
        """
        Validate artifacts returned by the
        multi-agent workflow.
        """

        artifact_paths = []

        # ------------------------------------------------------
        # DOCX
        # ------------------------------------------------------

        document = workflow_result.get(
            "generated_document"
        )

        if document:

            file_path = document.get(
                "file_path"
            )

            if file_path:

                artifact_paths.append(
                    file_path
                )

        # ------------------------------------------------------
        # PPTX
        # ------------------------------------------------------

        presentation = workflow_result.get(
            "generated_presentation"
        )

        if presentation:

            file_path = presentation.get(
                "file_path"
            )

            if file_path:

                artifact_paths.append(
                    file_path
                )

        # ------------------------------------------------------
        # Validate
        # ------------------------------------------------------

        if not artifact_paths:

            return {
                "valid": False,
                "artifact_count": 0,
                "artifacts": [],
                "errors": [
                    "No generated artifacts found."
                ],
            }

        return self.validate_artifacts(
            artifact_paths
        )